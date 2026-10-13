"""Zudo Works company deck: 8 slides, 16:9, native shapes and editable text. Called by make_brochure.py.

Same rules as the brochure: no prices, no client names or results, no partner claim; example workflows are
labelled as examples. Text uses Segoe UI (PowerPoint substitutes a similar sans-serif where it is missing).
"""
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

FONT = 'Segoe UI'
INK, BODY, MUTED, LINE, SOFT = '0F1B2D', '3D4A5C', '6B7689', 'E3E8EF', 'F4F7FB'
BLUE, BLUE_DK, BLUE_BG, GREEN, GREEN_BG = '226DB4', '1B5A96', 'EBF3FA', '089949', 'E6F7EE'
RED, AMBER, AMBER_BG, NAVY, NAVY2, SKY = 'E42527', 'F9B21D', 'FDF3E1', '0E1E36', '16294A', '7FB6EA'
W, H = 13.333, 7.5


def rgb(h): return RGBColor.from_string(h)


def box(s, x, y, w, h, fill=None, line=None, r=None, weight=0.75):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if r:
        shp.adjustments[0] = min(0.5, r / min(w, h))
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line); shp.line.width = Pt(weight)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    shp.text_frame.text = ''
    return shp


def oval(s, x, y, d, fill, line=None):
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line); shp.line.width = Pt(1.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def text(s, x, y, w, h, paras, size=12, color=BODY, bold=False, align='l', anchor='t', spacing=1.15, after=0):
    """paras: a string, or a list of paragraphs; each paragraph is a string or a list of (text, {opts}) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        p.line_spacing = spacing
        p.space_after = Pt(after)
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold)
            f.color.rgb = rgb(o.get('color', color))
    return tb


def pic(s, path, x, y, w=None, h=None):
    return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def pic_fill(s, path, x, y, w, h, focus=0.5):
    """Place an image to cover the box, cropping the overflow (like CSS object-fit: cover)."""
    iw, ih = Image.open(path).size
    sp = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    box_r, img_r = w / h, iw / ih
    if img_r > box_r:
        cut = 1 - box_r / img_r
        sp.crop_left, sp.crop_right = cut * focus, cut * (1 - focus)
    else:
        cut = 1 - img_r / box_r
        sp.crop_top, sp.crop_bottom = cut * focus, cut * (1 - focus)
    return sp


def fit(path, max_w, max_h):
    iw, ih = Image.open(path).size
    k = min(max_w / iw, max_h / ih)
    return iw * k, ih * k


def brand_bar(s, x, y, w=0.9, h=0.06):
    seg = (w - 3 * 0.04) / 4
    for i, c in enumerate((BLUE, GREEN, RED, AMBER)):
        box(s, x + i * (seg + 0.04), y, seg, h, fill=c, r=h / 2)


def chips(s, x, y, max_w, items, size=10, fill=SOFT, line=LINE, color=INK, h=0.32, gap=0.08):
    cx, cy = x, y
    for t in items:
        w = len(t) * size * 0.0078 + 0.32
        if cx + w > x + max_w:
            cx, cy = x, cy + h + gap
        b = box(s, cx, cy, w, h, fill=fill, line=line, r=h / 2)
        text(s, cx, cy, w, h, t, size=size, color=color, bold=True, align='c', anchor='m')
        cx += w + gap
    return cy + h


def frame(s, img, n, section):
    """Running head and foot for content slides."""
    pic(s, img['logo'], 0.6, 0.38, 0.3, 0.3)
    text(s, 0.98, 0.38, 3, 0.3, 'Zudo Works', size=11, color=INK, bold=True, anchor='m')
    text(s, W - 5.6, 0.38, 5, 0.3, section.upper(), size=9, color=MUTED, bold=True, align='r', anchor='m')
    box(s, 0.6, 0.86, W - 1.2, 0.012, fill=LINE)
    text(s, 0.6, H - 0.5, 4, 0.25, 'zudoworks.com', size=9, color=MUTED)
    text(s, W - 1.6, H - 0.5, 1, 0.25, f'{n:02d} / 08', size=9, color=INK, bold=True, align='r')


def heading(s, kicker, title, accent=None, y=1.15, size=30, w=11.5, color=INK, kcolor=BLUE):
    text(s, 0.6, y, w, 0.3, kicker.upper(), size=10, color=kcolor, bold=True)
    runs = [(title, {})]
    if accent:
        a, b = title.split(accent, 1)
        runs = [(a, {}), (accent, {'color': BLUE if color == INK else SKY}), (b, {})]
    text(s, 0.6, y + 0.32, w, 1.2, [runs], size=size, color=color, bold=True, spacing=1.0)


def build(path, img, sc):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    prs.core_properties.title = 'Zudo Works company overview'
    prs.core_properties.author = 'Zudo Works'
    prs.core_properties.subject = 'Zoho and custom software development'

    # 1 · Cover
    s = prs.slides.add_slide(blank)
    box(s, 0, 0, W, H, fill='F3F7FC')
    pic_fill(s, img['handshake'], 6.9, 0, W - 6.9, H, focus=0.5)
    box(s, 6.9, 0, 0.02, H, fill=LINE)
    pic(s, img['logo'], 0.7, 0.6, 0.55, 0.55)
    text(s, 1.38, 0.6, 4, 0.55, [[('Zudo Works', {'size': 18, 'bold': True, 'color': INK})], [('ZOHO & CUSTOM SOFTWARE DEVELOPMENT', {'size': 7.5, 'bold': True, 'color': MUTED})]], anchor='m', spacing=1.0)
    brand_bar(s, 0.7, 2.0)
    text(s, 0.7, 2.2, 6, 0.3, 'COMPANY OVERVIEW', size=10, color=BLUE, bold=True)
    text(s, 0.7, 2.55, 6.0, 2.1, [[('Zoho and custom software, built around ', {}), ('how your business actually runs', {'color': BLUE})]], size=30, color=INK, bold=True, spacing=0.98)
    text(s, 0.7, 4.8, 5.7, 0.9, 'We design, build and connect business systems on Zoho, and write the custom software when Zoho is not the right tool. You talk to the engineers who build it.', size=12.5, color=BODY, spacing=1.25)
    chips(s, 0.7, 5.85, 5.9, ['Zoho Creator', 'Zoho CRM', 'Deluge', 'Integrations', 'AI automation', 'Custom software'], size=9.5, fill='FFFFFF')
    text(s, 0.7, 6.75, 6, 0.4, [[('zudoworks.com', {'bold': True, 'color': INK, 'size': 12})], [('US · UK · Australia · New Zealand · Middle East · India', {'size': 9, 'color': MUTED})]], spacing=1.1)

    # 2 · Who we are
    s = prs.slides.add_slide(blank)
    frame(s, img, 2, 'Who we are')
    heading(s, 'Who we are', 'Engineers who build on Zoho every day', 'Zoho')
    text(s, 0.6, 2.05, 12, 0.7, 'A Zoho and custom software development company based in Chennai, India. We replace manual processes with automated workflows, connect systems that don’t talk to each other, and build custom software when an off-the-shelf product can’t do the job.', size=12.5, color=BODY, spacing=1.2)
    for i, (k, kc, t) in enumerate((('OUR MISSION', BLUE, 'To give growing businesses software that fits how they actually work: Zoho where it fits, custom code where it earns its place, built by the engineers they talk to, and owned by them.'),
                                    ('OUR VISION', GREEN, 'Every growing business running on one connected system, so its people spend their time on customers and craft instead of retyping and chasing.'))):
        x = 0.6 + i * 6.15
        box(s, x, 2.95, 5.95, 1.75, fill='FFFFFF', line=LINE, r=0.16)
        box(s, x + 0.25, 2.95, 5.45, 0.05, fill=kc)
        text(s, x + 0.35, 3.2, 5.3, 0.3, k, size=9.5, color=kc, bold=True)
        text(s, x + 0.35, 3.5, 5.3, 1.1, t, size=13, color=INK, spacing=1.2)
    stats = (('20+', 'Projects our team has worked on'), ('10+', 'Client businesses supported'), ('40+', 'Zoho products used'), ('1 month', 'Free support and onboarding after every build'))
    box(s, 0.6, 4.95, 12.1, 1.3, fill='FFFFFF', line=LINE, r=0.16)
    for i, (b, t) in enumerate(stats):
        x = 0.6 + i * 3.025
        if i:
            box(s, x, 5.1, 0.012, 1.0, fill=LINE)
        text(s, x + 0.3, 5.12, 2.7, 0.6, b, size=28, color=INK, bold=True)
        text(s, x + 0.3, 5.72, 2.6, 0.45, t, size=9.5, color=MUTED)
    text(s, 0.6, 6.38, 12, 0.3, 'Figures cover our team members’ project experience, much of it as developers at a previous organization.', size=8.5, color=MUTED)

    # 3 · Problems
    s = prs.slides.add_slide(blank)
    frame(s, img, 3, 'Where we help')
    heading(s, 'Where we help', 'The problems we are usually called in to fix', 'called in to fix')
    probs = (('Spreadsheets are running the business', 'Orders, approvals and job tracking live in shared sheets that break as soon as the team grows.'),
             ('The CRM is set up, but nobody trusts it', 'Duplicate records, skipped stages and reports that don’t match what the sales team says.'),
             ('Systems that don’t talk to each other', 'People re-key the same customer and order data between CRM, accounting and e-commerce.'),
             ('A Zoho app that grew without a plan', 'Creator apps and Deluge scripts that only one person understands, and nobody dares change.'),
             ('Follow-ups depend on someone remembering', 'Reminders, invoices, approvals and hand-offs that should happen automatically, but don’t.'),
             ('Off-the-shelf software almost fits', 'One critical workflow that no product supports without awkward workarounds.'))
    for i, (h, t) in enumerate(probs):
        x, y = 0.6 + (i % 3) * 4.07, 2.3 + (i // 3) * 2.0
        box(s, x, y, 3.9, 1.85, fill='FFFFFF', line=LINE, r=0.16)
        box(s, x + 0.3, y + 0.3, 0.42, 0.42, fill=AMBER_BG, r=0.1)
        text(s, x + 0.3, y + 0.3, 0.42, 0.42, str(i + 1), size=12, color='B07508', bold=True, align='c', anchor='m')
        text(s, x + 0.9, y + 0.3, 2.8, 0.42, h, size=12, color=INK, bold=True, anchor='m', spacing=1.0)
        text(s, x + 0.3, y + 0.95, 3.35, 0.8, t, size=10, color=MUTED, spacing=1.2)
    box(s, 0.6, 6.35, 12.1, 0.55, fill=NAVY, r=0.14)
    text(s, 0.9, 6.35, 11.6, 0.55, [[('We start from ', {}), ('how your business actually runs', {'color': SKY}), (', then choose the technology. Never the other way round.', {})]], size=13, color='FFFFFF', bold=True, anchor='m')

    # 4 · Services
    s = prs.slides.add_slide(blank)
    frame(s, img, 4, 'What we build')
    heading(s, 'What we build', 'Seven services, one connected system', 'one connected system')
    svcs = (('Zoho Creator development', 'Custom apps: data models, forms, client portals, dashboards and Deluge workflows, documented so your team can own them.'),
            ('Zoho CRM implementation', 'Pipelines and blueprints that match how you sell, custom modules, data migration and training.'),
            ('Deluge development', 'Functions, scheduled jobs, buttons and API calls across CRM, Creator and Books, plus reviews and fixes.'),
            ('Zoho integrations', 'Connect Zoho with accounting, e-commerce, payments and your own systems via APIs, webhooks, Zoho Flow or middleware.'),
            ('Business process automation', 'We map the process first, then remove the manual steps across sales, finance and operations.'),
            ('Custom software', 'Web apps, internal tools and dashboards where Zoho is the wrong tool, built to integrate with what you run.'))
    for i, (h, t) in enumerate(svcs):
        x, y = 0.6 + (i % 3) * 4.07, 2.15 + (i // 3) * 1.85
        box(s, x, y, 3.9, 1.65, fill='FFFFFF', line=LINE, r=0.16)
        box(s, x + 0.3, y + 0.3, 0.42, 0.42, fill=BLUE_BG, r=0.1)
        text(s, x + 0.3, y + 0.3, 0.42, 0.42, f'{i + 1:02d}', size=10, color=BLUE, bold=True, align='c', anchor='m')
        text(s, x + 0.85, y + 0.3, 2.9, 0.42, h, size=12.5, color=INK, bold=True, anchor='m')
        text(s, x + 0.3, y + 0.85, 3.35, 0.75, t, size=9.5, color=MUTED, spacing=1.15)
    box(s, 0.6, 5.95, 12.1, 0.95, fill=BLUE_BG, line='CFE0F1', r=0.16)
    text(s, 0.9, 5.95, 3.2, 0.95, '07  Support and maintenance', size=12.5, color=BLUE_DK, bold=True, anchor='m')
    text(s, 4.1, 5.95, 8.4, 0.95, 'Every build includes one month of free support and onboarding after launch. After that, optional monthly support: fixes, small changes, new users and Zoho updates.', size=10.5, color=BODY, anchor='m', spacing=1.15)

    # 5 · Connected
    s = prs.slides.add_slide(blank)
    frame(s, img, 5, 'How it connects')
    heading(s, 'How it connects', 'From a tangle of tools to one integration and AI layer', 'one integration and AI layer', size=26, w=12)
    ew, eh = fit(img['explorer'], 5.6, 4.15)
    box(s, 0.6 - 0.06, 2.1 - 0.06, ew + 0.12, eh + 0.12, fill=SOFT, line=LINE, r=0.2)
    pic(s, img['explorer'], 0.6, 2.1, ew, eh)
    layers = (('1  Untangle', 'We map every app, sheet and hand-off, and decide which system owns which record.'),
              ('2  Connect', 'Deluge, Zoho Flow and APIs link each app to the layer once, instead of to every other app.'),
              ('3  Add AI where it is checkable', 'Zoho Zia, or Claude and ChatGPT through Zoho MCP, read documents and draft records. A person approves.'),
              ('4  Sync back', 'Status, invoices and payments flow back to the record that started the work.'))
    x0 = 0.6 + ew + 0.5
    for i, (h, t) in enumerate(layers):
        y = 2.1 + i * 0.86
        if i:
            box(s, x0, y - 0.1, W - 0.6 - x0, 0.012, fill=LINE)
        text(s, x0, y, W - 0.6 - x0, 0.3, h, size=12.5, color=INK, bold=True)
        text(s, x0, y + 0.3, W - 0.6 - x0, 0.5, t, size=9.5, color=MUTED, spacing=1.15)
    lw, lh = fit(img['logos'], W - 0.6 - x0, 1.2)
    pic(s, img['logos'], x0, 5.65, lw, lh)
    text(s, 0.6, 2.1 + eh + 0.12, ew, 0.3, 'Example workflow from our interactive explorer. Not a client case study.', size=8.5, color=MUTED)

    # 6 · How we work
    s = prs.slides.add_slide(blank)
    frame(s, img, 6, 'How a project runs')
    heading(s, 'How a project runs', 'From first call to a working system', 'working system')
    steps = (('Understand first', 'We follow real work from start to finish and note every retyping and hand-off.'),
             ('Agree a first step', 'A written proposal with scope, timeline and cost before any work starts.'),
             ('Build and trial', 'Built in stages with regular demos, changed in response to the people using it.'),
             ('Launch and stay', 'Training, documentation and one month of free support after launch.'))
    box(s, 0.95, 2.48, 9.2, 0.025, fill=LINE)
    for i, (h, t) in enumerate(steps):
        x = 0.6 + i * 3.07
        last = i == 3
        oval(s, x, 2.15, 0.68, GREEN if last else 'FFFFFF', None if last else BLUE)
        text(s, x, 2.15, 0.68, 0.68, str(i + 1), size=15, color='FFFFFF' if last else BLUE, bold=True, align='c', anchor='m')
        text(s, x, 3.0, 2.8, 0.35, h, size=13, color=INK, bold=True)
        text(s, x, 3.35, 2.8, 0.7, t, size=9.5, color=MUTED, spacing=1.15)
    box(s, 0.6, 4.3, 5.4, 2.6, fill=SOFT, r=0.16)
    text(s, 0.9, 4.48, 5, 0.35, 'Every project includes', size=14, color=INK, bold=True)
    incl = ['A written proposal before any work starts', 'Design, build and testing in stages, with demos', 'Documentation your team can maintain',
            'Training for administrators and users', 'One month of free support after launch', 'Zoho licences and data in your own name']
    for i, t in enumerate(incl):
        y = 4.95 + i * 0.31
        oval(s, 0.92, y + 0.08, 0.12, GREEN)
        text(s, 1.2, y, 4.6, 0.28, t, size=10, color=INK)
    lw, lh = fit(img['lanes'], 6.5, 2.6)
    pic(s, img['lanes'], W - 0.6 - lw, 4.3 + (2.6 - lh) / 2, lw, lh)

    # 7 · Industries and manufacturing
    s = prs.slides.add_slide(blank)
    frame(s, img, 7, 'Solutions by industry')
    heading(s, 'Solutions by industry', '90 example workflows across 18 industries', '90 example workflows', size=26, w=12)
    text(s, 0.6, 2.05, 5.2, 0.3, 'INDUSTRIES', size=9, color=MUTED, bold=True)
    end = chips(s, 0.6, 2.35, 5.3, [i['name'] for i in sc.INDUSTRIES], size=9)
    text(s, 0.6, end + 0.2, 5.2, 0.9, 'Five common problems per industry, each solved on Zoho with integrations and AI. Explore them all at zudoworks.com/solutions. Example workflows, not client case studies.', size=9.5, color=MUTED, spacing=1.2)
    x0 = 6.3
    box(s, x0, 2.0, W - 0.6 - x0, 4.9, fill=SOFT, r=0.18)
    text(s, x0 + 0.35, 2.2, 5, 0.3, 'SPOTLIGHT · MANUFACTURING', size=9, color=BLUE, bold=True)
    text(s, x0 + 0.35, 2.48, 6, 0.5, 'Manufacturing software built around how your factory runs', size=14, color=INK, bold=True)
    ow, oh = fit(img['ops'], W - 0.6 - x0 - 0.7, 2.35)
    pic(s, img['ops'], x0 + 0.35, 3.1, ow, oh)
    mods = ['Job and production tracking', 'Inventory and purchasing', 'Quality, NCR and traceability', 'Dashboards and AI order intake']
    for i, t in enumerate(mods):
        x, y = x0 + 0.35 + (i % 2) * 3.0, 3.1 + oh + 0.2 + (i // 2) * 0.32
        oval(s, x + 0.02, y + 0.08, 0.12, GREEN)
        text(s, x + 0.28, y, 2.7, 0.28, t, size=10, color=INK)

    # 8 · Contact
    s = prs.slides.add_slide(blank)
    box(s, 0, 0, W, H, fill=NAVY)
    pic(s, img['logo'], 0.6, 0.5, 0.42, 0.42)
    text(s, 1.12, 0.5, 3, 0.42, 'Zudo Works', size=14, color='FFFFFF', bold=True, anchor='m')
    text(s, W - 3.6, 0.5, 3, 0.42, '08 / 08', size=9, color='8597B2', bold=True, align='r', anchor='m')
    heading(s, 'Get started', 'Tell us how your business runs today', None, y=1.3, size=32, color='FFFFFF', kcolor=SKY)
    text(s, 0.6, 2.25, 7.2, 0.7, 'Describe one process from start to finish. We will tell you what we would build first, what we would leave alone, and how long it would take.', size=12.5, color='C9D5E6', spacing=1.2)
    for i, (h, t) in enumerate((('We reply within 6 hours, guaranteed', 'With a few questions or times for a call.'),
                                 ('A discovery call', 'About how the work flows today.'),
                                 ('A written proposal', 'Scope, timeline and cost, before work starts.'))):
        x = 0.6 + i * 2.5
        box(s, x, 3.25, 2.35, 1.45, fill=NAVY2, r=0.14)
        oval(s, x + 0.25, 3.45, 0.36, GREEN)
        text(s, x + 0.25, 3.45, 0.36, 0.36, str(i + 1), size=11, color='FFFFFF', bold=True, align='c', anchor='m')
        text(s, x + 0.25, 3.9, 1.95, 0.45, h, size=10.5, color='FFFFFF', bold=True, spacing=1.05)
        text(s, x + 0.25, 4.3, 1.95, 0.4, t, size=8.5, color='C9D5E6')
    box(s, 0.6, 5.0, 7.35, 1.85, fill='FFFFFF', r=0.18)
    rows = (('WEBSITE', 'zudoworks.com'), ('EMAIL', 'admin@zudoworks.com'), ('BOOK A CALL', 'zudoworks.com/contact'), ('BASED IN', 'Anna Nagar, Chennai, India'))
    for i, (k, v) in enumerate(rows):
        y = 5.2 + i * 0.38
        text(s, 0.95, y, 1.5, 0.32, k, size=8.5, color=MUTED, bold=True, anchor='m')
        text(s, 2.45, y, 3.6, 0.32, v, size=12, color=INK, bold=True, anchor='m')
    pic(s, img['qr'], 6.3, 5.15, 1.35, 1.35)
    text(s, 6.1, 6.52, 1.75, 0.25, 'Scan to send details', size=7.5, color=MUTED, align='c')
    # Who they will talk to
    x0 = 8.45
    box(s, x0, 1.3, W - 0.6 - x0, 5.55, fill=NAVY2, r=0.2)
    pic_fill(s, img['cto'], x0 + 0.35, 1.65, W - 0.6 - x0 - 0.7, 2.7, focus=0.3)
    text(s, x0 + 0.35, 4.5, 4, 0.35, 'Arunkumar V', size=16, color='FFFFFF', bold=True)
    text(s, x0 + 0.35, 4.85, 4, 0.3, 'Chief Technology Officer', size=10, color=SKY, bold=True)
    text(s, x0 + 0.35, 5.25, W - 0.6 - x0 - 0.7, 0.7, '“Master of Creator (Global Winner)” award, Zoho Creator Partner Hackathon 2025', size=10, color='FFFFFF', spacing=1.15)
    text(s, x0 + 0.35, 5.95, W - 0.6 - x0 - 0.7, 0.7, 'Makes the technical decisions on every project. You talk to the people who build it.', size=9, color='C9D5E6', spacing=1.15)
    text(s, 0.6, H - 0.42, 10, 0.25, 'zudoworks.com is the official and only website of Zudo Works.', size=8, color='8597B2')

    prs.save(path)
    return len(prs.slides)
